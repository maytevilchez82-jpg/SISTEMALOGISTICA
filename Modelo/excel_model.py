import os
import re
import math
import unicodedata
from datetime import datetime
from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile
from threading import Lock

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from pyxlsb import open_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_EXCEL_PATH = BASE_DIR / "inventario.xlsx"
SOURCE_XLSB_PATH = BASE_DIR / "INVENTARIO DE EQUIPOS ENERO 2026.xlsb"


def _resolve_excel_path():
    env_path = os.getenv("EXCEL_PATH") or os.getenv("DB_PATH")
    if env_path:
        return Path(env_path).expanduser().resolve()

    custom_excel = (
        sorted(BASE_DIR.glob("*.xlsb"))
        + sorted(BASE_DIR.glob("*.xlsx"))
        + sorted(BASE_DIR.glob("*.xlsm"))
        + sorted(BASE_DIR.glob("*.xls"))

    )
    if custom_excel:
        return custom_excel[0]

    return DEFAULT_EXCEL_PATH


SOURCE_EXCEL_PATH = _resolve_excel_path()
EXCEL_PATH = DEFAULT_EXCEL_PATH if SOURCE_EXCEL_PATH.suffix.lower() == ".xlsb" else SOURCE_EXCEL_PATH
excel_lock = Lock()

PRODUCT_HEADERS = ["id", "sku", "name", "description", "brand", "model", "part_number", "project", "unit", "group", "category", "floor", "location", "position", "section", "observation", "source_sheet", "item", "stock", "min_stock", "price", "updated_at"]
MOVEMENT_HEADERS = ["id", "product_id", "movement_type", "quantity", "floor", "location", "position", "project", "note", "created_at", "group"]


def _parse_excel_number(value, field_name, sheet_name, row_number=None):
    if value is None or isinstance(value, str) and value.strip().casefold() in {"", "-", "—", "n/a", "n/d", "na", "s/d", "sin dato", "sin stock", "no aplica"}:
        return 0.0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
    else:
        text = unicodedata.normalize("NFKC", str(value)).strip().replace("−", "-")
        negative = text.startswith("(") and text.endswith(")")
        if negative:
            text = text[1:-1]
        text = re.sub(r"(?i)US\$|USD|PEN|S/", "", text)
        text = re.sub(r"[\s$€£¥A-Za-z]", "", text)
        if text in {"", "-", "—"}:
            return 0.0
        if "," in text and "." in text:
            decimal_separator = "," if text.rfind(",") > text.rfind(".") else "."
            thousands_separator = "." if decimal_separator == "," else ","
            text = text.replace(thousands_separator, "")
            if decimal_separator == ",":
                text = text.replace(",", ".")
        elif "," in text:
            if text.count(",") > 1:
                pieces = text.split(",")
                text = "".join(pieces) if len(pieces[-1]) == 3 else "".join(pieces[:-1]) + "." + pieces[-1]
            else:
                whole, fraction = text.split(",")
                text = whole + fraction if len(fraction) == 3 else whole + "." + fraction
        elif text.count(".") > 1:
            pieces = text.split(".")
            text = "".join(pieces) if len(pieces[-1]) == 3 else "".join(pieces[:-1]) + "." + pieces[-1]
        elif "." in text and len(text.rsplit(".", 1)[1]) == 3:
            text = text.replace(".", "")
        try:
            number = float(text)
            if negative:
                number = -number
        except ValueError as error:
            location = f" en la hoja '{sheet_name}'" + (f", fila {row_number}" if row_number else "")
            raise ValueError(f"El campo {field_name}{location} contiene un valor no numérico.") from error
    if not math.isfinite(number):
        location = f" en la hoja '{sheet_name}'" + (f", fila {row_number}" if row_number else "")
        raise ValueError(f"El campo {field_name}{location} contiene un valor no numérico.")
    return number


def _rows_from_sheet(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row)]


def read_data():
    with excel_lock:
        workbook = load_workbook(EXCEL_PATH, data_only=True)
        products = _rows_from_sheet(workbook["Productos"])
        movements = _rows_from_sheet(workbook["Movimientos"])
    for product in products:
        product.setdefault("unit", "Unidades")
        product.setdefault("description", product.get("name", ""))
        product.setdefault("brand", "")
        product.setdefault("model", "")
        product.setdefault("part_number", "")
        product.setdefault("project", "")
        product.setdefault("group", product.get("category", "Consumibles"))
        product.setdefault("category", product.get("group", "Consumibles"))
        product.setdefault("floor", "")
        product.setdefault("location", "")
        product.setdefault("position", "")
        product.setdefault("section", "")
        product.setdefault("observation", "")
        product.setdefault("source_sheet", "")
        product.setdefault("item", product.get("id", ""))
    return products, movements


def write_data(products, movements):
    workbook = Workbook()
    products_sheet = workbook.active
    products_sheet.title = "Productos"
    movements_sheet = workbook.create_sheet("Movimientos")
    products_sheet.append(PRODUCT_HEADERS)
    movements_sheet.append(MOVEMENT_HEADERS)
    for product in products:
        products_sheet.append([product.get(header) for header in PRODUCT_HEADERS])
    for movement in movements:
        movements_sheet.append([movement.get(header) for header in MOVEMENT_HEADERS])
    for sheet in (products_sheet, movements_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    products_sheet.column_dimensions["B"].width = 16
    products_sheet.column_dimensions["C"].width = 28
    products_sheet.column_dimensions["D"].width = 18
    movements_sheet.column_dimensions["E"].width = 28
    with excel_lock:
        workbook.save(EXCEL_PATH)


def initialize_database():
    if EXCEL_PATH.exists():
        return
    if SOURCE_EXCEL_PATH.exists() and SOURCE_EXCEL_PATH.suffix.lower() == ".xlsb":
        write_data(_import_xlsb_products(SOURCE_EXCEL_PATH), [])
        return
    if SOURCE_XLSB_PATH.exists():
        write_data(_import_xlsb_products(SOURCE_XLSB_PATH), [])
        return
    now = datetime.now().isoformat(timespec="seconds")
    products = [
        {"id": 1, "sku": "TEC-001", "name": "Teclado mecanico", "unit": "Cajas", "group": "Perifericos", "category": "Perifericos", "stock": 24, "min_stock": 8, "price": 49.90, "updated_at": now},
        {"id": 2, "sku": "MON-024", "name": "Monitor 24 pulgadas", "unit": "Cajas", "group": "Monitores", "category": "Monitores", "stock": 8, "min_stock": 4, "price": 189.00, "updated_at": now},
        {"id": 3, "sku": "CAB-110", "name": "Cable USB-C reforzado", "unit": "Cajas", "group": "Accesorios", "category": "Accesorios", "stock": 42, "min_stock": 12, "price": 14.50, "updated_at": now},
        {"id": 4, "sku": "AUD-301", "name": "Audifonos estudio", "unit": "Cajas", "group": "Audio", "category": "Audio", "stock": 3, "min_stock": 5, "price": 79.90, "updated_at": now},
        {"id": 5, "sku": "MOU-089", "name": "Mouse ergonomico", "unit": "Cajas", "group": "Perifericos", "category": "Perifericos", "stock": 17, "min_stock": 6, "price": 32.00, "updated_at": now},
    ]
    movements = [
        {"id": 1, "product_id": 1, "movement_type": "Entrada", "quantity": 10, "note": "Compra proveedor", "created_at": now},
        {"id": 2, "product_id": 4, "movement_type": "Salida", "quantity": 2, "note": "Pedido #1048", "created_at": now},
        {"id": 3, "product_id": 3, "movement_type": "Entrada", "quantity": 20, "note": "Reposicion", "created_at": now},
        {"id": 4, "product_id": 2, "movement_type": "Salida", "quantity": 1, "note": "Pedido #1047", "created_at": now},
    ]
    write_data(products, movements)


def _import_xlsb_products(source_path=None):
    source_path = Path(source_path) if source_path else SOURCE_XLSB_PATH
    products = []
    next_id = 1
    with open_workbook(source_path) as workbook:
        for sheet_name in workbook.sheets:
            with workbook.get_sheet(sheet_name) as sheet:
                rows = ([cell.v for cell in row] for row in sheet.rows())
                headers = None
                for row in rows:
                    normalized = [str(value).strip().upper() if value is not None else "" for value in row]
                    if "ITEM" in normalized and "DESCRIPCIÓN" in normalized:
                        headers = normalized
                        break
                if not headers:
                    continue
                for row in rows:
                    values = list(row)
                    if not any(value is not None for value in values):
                        continue
                    data = {headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))}
                    description = str(data.get("DESCRIPCIÓN") or "").strip()
                    if not description:
                        continue
                    item = data.get("ITEM")
                    item_text = str(int(item)) if isinstance(item, float) and item.is_integer() else str(item or "")
                    quantity = int(_parse_excel_number(data.get("CANTIDAD"), "cantidad", sheet_name))
                    location = data.get("CASILLERO")
                    location_text = f"Locker {int(location)}" if isinstance(location, float) and location.is_integer() else str(location or "").strip()
                    products.append({
                        "id": next_id,
                        "sku": f"{sheet_name}-{item_text}",
                        "name": description,
                        "description": description,
                        "brand": str(data.get("MARCA") or "").strip(),
                        "model": str(data.get("MODELO") or "").strip(),
                        "part_number": str(data.get("PART NUMBER") or "").strip(),
                        "unit": "Unidades",
                        "group": sheet_name,
                        "category": sheet_name,
                        "floor": "",
                        "location": location_text,
                        "position": str(data.get("POSICION") or "").strip(),
                        "observation": str(data.get("OBSERVACION") or "").strip(),
                        "source_sheet": sheet_name,
                        "item": item_text,
                        "stock": quantity,
                        "min_stock": 0,
                        "price": 0,
                        "updated_at": datetime.now().isoformat(timespec="seconds"),
                    })
                    next_id += 1
    return products


def get_products():
    products, _ = read_data()
    unique_products = {}
    for product in products:
        key = _product_identity(product)
        if key not in unique_products:
            unique_products[key] = dict(product)
        else:
            unique_products[key]["stock"] = int(unique_products[key].get("stock") or 0) + int(product.get("stock") or 0)
    return sorted(unique_products.values(), key=lambda item: item["name"])


def _product_identity(product):
    return _movement_product_identity(product) + (_product_inventory_group(product),)


def _sync_imported_products(imported):
    existing_products, movements = read_data()
    products_by_sku = {str(product.get("sku", "")).strip().upper(): product for product in existing_products}
    next_id = max((int(product.get("id") or 0) for product in existing_products), default=0) + 1
    synchronized_products = []
    synchronized_skus = set()

    for product in imported:
        sku = str(product.get("sku", "")).strip().upper()
        existing = products_by_sku.get(sku)
        if existing:
            existing.update(product)
            if sku not in synchronized_skus:
                synchronized_products.append(existing)
                synchronized_skus.add(sku)
        else:
            product["id"] = next_id
            next_id += 1
            synchronized_products.append(product)
            products_by_sku[sku] = product
            synchronized_skus.add(sku)

    write_data(synchronized_products, movements)


def import_products_workbook(uploaded_file):
    """Synchronize uploaded workbook sheets with the stored inventory."""
    if str(getattr(uploaded_file, "filename", "")).lower().endswith(".xlsb"):
        with NamedTemporaryFile(suffix=".xlsb", delete=False) as temporary_file:
            temporary_path = Path(temporary_file.name)
            uploaded_file.save(temporary_file)
        try:
            imported = _import_xlsb_products(temporary_path)
        finally:
            temporary_path.unlink(missing_ok=True)
        if not imported:
            raise ValueError("No se encontraron filas de inventario con encabezados reconocibles.")
        _sync_imported_products(imported)
        return len(imported)

    aliases = {
        "sku": {"sku", "codigo", "codigoproducto"},
        "item": {"item", "nroitem", "numeroitem"},
        "name": {"descripcion", "producto", "nombre", "articulo", "description", "name"},
        "brand": {"marca", "brand"},
        "model": {"modelo", "model"},
        "part_number": {"partnumber", "numeroparte"},
        "project": {"proyecto", "project"},
        "unit": {"unidad", "unidadmedida"},
        "group": {"grupo", "tipo", "hoja"},
        "category": {"categoria"},
        "floor": {"piso"},
        "location": {"casillero", "locker", "locacion", "ubicacion"},
        "position": {"posicion", "position"},
        "section": {"seccion", "section"},
        "observation": {"observacion", "observaciones", "nota"},
        "stock": {"cantidad", "existencia", "stock", "qty"},
        "min_stock": {"stockminimo", "minimo", "minstock"},
        "price": {"precio", "preciounitario"},
    }

    def normalize(value):
        text = unicodedata.normalize("NFKD", str(value or ""))
        text = "".join(char for char in text if not unicodedata.combining(char))
        return re.sub(r"[^a-z0-9]", "", text.lower())

    workbook = load_workbook(uploaded_file, data_only=True, read_only=True)
    imported = []
    for sheet in workbook.worksheets:
        header_row = None
        field_columns = {}
        for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            normalized = [normalize(value) for value in row]
            columns = {
                field: next((index for index, value in enumerate(normalized) if value in names), None)
                for field, names in aliases.items()
            }
            if columns["name"] is not None:
                header_row = row_index
                field_columns = columns
                break
        if header_row is None:
            continue

        def value_at(row, field, default=""):
            index = field_columns.get(field)
            return row[index] if index is not None and index < len(row) and row[index] is not None else default

        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=1):
            name = str(value_at(row, "name")).strip()
            if not name:
                continue
            excel_row = header_row + row_number
            item = str(value_at(row, "item", row_number)).strip()
            row_group = str(value_at(row, "group") or sheet.title).strip()
            sku = str(value_at(row, "sku") or f"{row_group}-{item}").strip().upper()
            if not sku or sku == f"{row_group}-":
                sku = f"{row_group}-{row_number}"
            category = str(value_at(row, "category") or row_group).strip()
            imported.append({
                "sku": sku,
                "name": name,
                "description": name,
                "brand": str(value_at(row, "brand")).strip(),
                "model": str(value_at(row, "model")).strip(),
                "part_number": str(value_at(row, "part_number")).strip(),
                "project": str(value_at(row, "project")).strip(),
                "unit": str(value_at(row, "unit") or "Unidades").strip(),
                "group": row_group,
                "category": category,
                "floor": str(value_at(row, "floor")).strip(),
                "location": str(value_at(row, "location")).strip(),
                "position": str(value_at(row, "position")).strip(),
                "section": str(value_at(row, "section")).strip(),
                "observation": str(value_at(row, "observation")).strip(),
                "source_sheet": row_group,
                "item": item,
                "stock": int(_parse_excel_number(value_at(row, "stock", 0), "cantidad", sheet.title, excel_row)),
                "min_stock": int(_parse_excel_number(value_at(row, "min_stock", 0), "stock mínimo", sheet.title, excel_row)),
                "price": _parse_excel_number(value_at(row, "price", 0), "precio", sheet.title, excel_row),
                "updated_at": datetime.now().isoformat(timespec="seconds"),
            })
    workbook.close()
    if not imported:
        raise ValueError("No se encontraron filas de inventario con encabezados reconocibles.")

    _sync_imported_products(imported)
    return len(imported)


def _fit_table_columns(sheet, min_row, max_row, max_column, width_limits=None):
    width_limits = width_limits or {}
    column_widths = {}
    for column_index in range(1, max_column + 1):
        longest = max(
            (len(line) for row_index in range(min_row, max_row + 1)
             for line in str(sheet.cell(row=row_index, column=column_index).value or "").splitlines()),
            default=0,
        )
        width = min(max(longest + 2, 12), width_limits.get(column_index, 255))
        column_widths[column_index] = width
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index in range(min_row, max_row + 1):
        line_count = 1
        for column_index in range(1, max_column + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.alignment = Alignment(
                horizontal="center" if row_index == min_row else cell.alignment.horizontal,
                vertical=cell.alignment.vertical,
                text_rotation=cell.alignment.text_rotation,
                wrap_text=True,
                shrink_to_fit=cell.alignment.shrink_to_fit,
                indent=cell.alignment.indent,
            )
            text = str(cell.value or "")
            width = max(column_widths[column_index] - 2, 1)
            wrapped_lines = sum(max(1, (len(line) + width - 1) // width) for line in text.splitlines())
            line_count = max(line_count, wrapped_lines)
        if line_count > 1:
            sheet.row_dimensions[row_index].height = min(line_count * 15, 409.5)


def export_products_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventario"
    logo_path = BASE_DIR / "static" / "NAKAMA LOGO 1.jpeg"
    if logo_path.exists():
        logo = Image(logo_path)
        logo.width = 240
        logo.height = 90
        sheet.add_image(logo, "A1")
    sheet.merge_cells("B1:N2")
    title = sheet["B1"]
    title.value = "INVENTARIO DE NAKAMA SOLUCIONES"
    title.font = Font(name="Calibri", size=18, bold=True, color="18236F")
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells("B3:N3")
    download_date = sheet["B3"]
    download_date.value = f"Fecha de descarga: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    download_date.font = Font(name="Calibri", size=11, italic=True, color="5C747C")
    download_date.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 40
    sheet.row_dimensions[2].height = 40
    sheet.row_dimensions[3].height = 10

    headers = ["ID", "TIPO", "ITEM", "PRODUCTO", "MARCA", "MODELO", "PART NUMBER", "UNIDAD", "CANTIDAD", "PISO", "LOCKER", "POSICION", "OBSERVACION", "ESTADO"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=4, column=column, value=header)
    for product in get_products():
        sheet.append([
            product["id"],
            product.get("source_sheet", ""),
            product.get("item", ""),
            product["name"],
            product.get("brand", ""),
            product.get("model", ""),
            product.get("part_number", ""),
            product.get("unit", "Unidades"),
            product["stock"],
            product.get("floor", ""),
            product.get("location", ""),
            product.get("position", ""),
            product.get("observation", ""),
            "Reponer" if product["stock"] <= product["min_stock"] else "Saludable",
        ])

    last_row = max(sheet.max_row, 5)
    table_border = Border(left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"), top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"))
    for row in sheet.iter_rows(min_row=4, max_row=last_row, min_col=1, max_col=14):
        for cell in row:
            cell.border = table_border
    table = Table(displayName="InventarioTabla", ref=f"A4:N{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "A5"
    _fit_table_columns(sheet, 4, last_row, 14, {4: 83})

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _movement_datetime(value):
    if isinstance(value, datetime):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        for date_format in ("%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(text[:10], date_format)
            except ValueError:
                continue
    return None


def export_movements_workbook(month=None, year=None):
    movements = get_movements()
    if month is not None or year is not None:
        movements = [
            movement for movement in movements
            if (movement_date := _movement_datetime(movement.get("created_at"))) is not None
            and (month is None or movement_date.month == month)
            and (year is None or movement_date.year == year)
        ]
    month_names = ["", "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
    period = " - ".join(value for value in (month_names[month] if month else "", str(year) if year else "") if value)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Entradas y salidas"
    logo_path = BASE_DIR / "static" / "NAKAMA LOGO 1.jpeg"
    if logo_path.exists():
        logo = Image(logo_path)
        logo.width = 200
        logo.height = 75
        sheet.add_image(logo, "A1")
    sheet.merge_cells("B1:H2")
    title = sheet["B1"]
    title.value = f"ENTRADAS Y SALIDAS DE STOCK{f' - {period}' if period else ''}"
    title.font = Font(name="Calibri", size=18, bold=True, color="18236F")
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells("B3:H3")
    download_date = sheet["B3"]
    download_date.value = f"Fecha de descarga: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    download_date.font = Font(name="Calibri", size=11, italic=True, color="5C747C")
    download_date.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 40
    sheet.row_dimensions[2].height = 40
    headers = ["PRODUCTO", "MARCA", "TIPO", "CANTIDAD", "FECHA", "PROYECTO", "OBSERVACIÓN"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=5, column=column, value=header)
    for movement in movements:
        sheet.append([
            movement.get("product_name", ""),
            movement.get("brand", ""),
            movement.get("movement_type", ""),
            movement.get("quantity", 0),
            movement.get("created_at", ""),
            movement.get("project", ""),
            movement.get("note", ""),
        ])
    last_row = max(sheet.max_row, 6)
    table_border = Border(left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"), top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"))
    for row in sheet.iter_rows(min_row=5, max_row=last_row, min_col=1, max_col=7):
        for cell in row:
            cell.border = table_border
            if cell.row == 5:
                cell.fill = PatternFill(fill_type="solid", fgColor="EAF5F7")
                cell.font = Font(name="Calibri", size=10, bold=True, color="31596B")
                cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[5].height = 24
    table = Table(displayName="MovimientosTabla", ref=f"A5:G{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    sheet.freeze_panes = "A6"
    _fit_table_columns(sheet, 5, last_row, 7)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def get_movements():
    products, movements = read_data()
    products_by_id = {product["id"]: product for product in products}
    return [movement_with_product(movement, products_by_id) for movement in sorted(movements, key=lambda item: item["id"], reverse=True)]


def movement_with_product(movement, products_by_id):
    product = products_by_id.get(movement["product_id"], {})
    inventory_group = product.get("group") or product.get("category") or product.get("source_sheet") or product.get("observation", "")
    return {**movement, "product_name": product.get("name", "Producto eliminado"), "sku": product.get("sku", ""), "model": product.get("model", ""), "part_number": product.get("part_number", ""), "brand": product.get("brand", ""), "project": movement.get("project") or product.get("project", ""), "inventory_group": inventory_group}


def _normalize_movement_group(value):
    normalized = unicodedata.normalize("NFD", str(value or ""))
    normalized = "".join(char for char in normalized if unicodedata.category(char) != "Mn")
    normalized = re.sub(r"[^a-z0-9]+", "", normalized.casefold())
    return {
        "ventas": "VENTA",
        "venta": "VENTA",
        "proyectos": "PROYECTO",
        "proyecto": "PROYECTO",
        "activos": "ACTIVO",
        "activo": "ACTIVO",
        "servicio": "SERVICIO",
        "servicios": "SERVICIO",
    }.get(normalized, str(value or "").strip().upper())


def _movement_product_identity(product):
    return tuple(
        re.sub(
            r"[^a-z0-9]+",
            " ",
            "".join(
                char for char in unicodedata.normalize("NFD", str(value or "").casefold())
                if unicodedata.category(char) != "Mn"
            ),
        ).strip()
        for value in (
            product.get("name") or product.get("description"),
            product.get("brand"),
            product.get("model"),
            product.get("part_number"),
        )
    )


def _product_inventory_group(product):
    for field in ("group", "category", "source_sheet"):
        value = product.get(field)
        if value:
            return _normalize_movement_group(value)
    return ""


def _new_group_product(products, source, group):
    product = dict(source)
    product_id = max((int(item.get("id") or 0) for item in products), default=0) + 1
    sku_base = f"{str(source.get('sku') or 'PRODUCTO').strip().upper()}-{group}"
    sku = sku_base
    suffix = 2
    existing_skus = {str(item.get("sku") or "").strip().upper() for item in products}
    while sku in existing_skus:
        sku = f"{sku_base}-{suffix}"
        suffix += 1
    product.update({
        "id": product_id,
        "sku": sku,
        "group": group,
        "category": group,
        "source_sheet": group,
        "item": str(product_id),
        "stock": 0,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    })
    products.append(product)
    return product


def _resolve_movement_product(products, product_id, group, movement_type):
    selected_product = next((item for item in products if int(item["id"]) == product_id), None)
    if not selected_product:
        return None
    identity = _movement_product_identity(selected_product)
    group_product = next((
        item for item in products
        if _movement_product_identity(item) == identity and _product_inventory_group(item) == group
    ), None)
    if group_product:
        return group_product
    if movement_type == "Entrada":
        return _new_group_product(products, selected_product, group)
    return None


def add_product(data):
    products, movements = read_data()
    sku = data["sku"].strip().upper()
    if any(product["sku"] == sku for product in products):
        return False, "El SKU ya existe."
    products.append({
        "id": max((product["id"] for product in products), default=0) + 1,
        "sku": sku,
        "name": data["name"].strip(),
        "description": data["name"].strip(),
        "brand": str(data.get("brand", "")).strip(),
        "model": str(data.get("model", "")).strip(),
        "part_number": str(data.get("part_number", "")).strip(),
        "project": str(data.get("project", "")).strip(),
        "unit": data.get("unit", "Unidades").strip(),
        "group": data["group"].strip(),
        "category": data["group"].strip(),
        "floor": str(data.get("floor", "Piso 1")).strip(),
        "location": str(data.get("location", "Locker 1")).strip(),
        "position": str(data.get("position", "")).strip(),
        "section": str(data.get("section", "")).strip(),
        "stock": int(data.get("stock", 0)),
        "min_stock": int(data.get("min_stock", 5)),
        "price": float(data.get("price", 0)),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    })
    write_data(products, movements)
    return True, "Producto creado correctamente."


def update_product(product_id, data):
    products, movements = read_data()
    product = next((item for item in products if int(item["id"]) == int(product_id)), None)
    if not product:
        return False, "Producto no encontrado."

    sku = str(data.get("sku", product.get("sku", ""))).strip().upper()
    if not sku:
        raise ValueError("El codigo es obligatorio.")
    if any(item is not product and str(item.get("sku", "")).strip().upper() == sku for item in products):
        return False, "El codigo ya pertenece a otro producto."

    text_fields = ("name", "description", "brand", "model", "part_number", "project", "unit", "group", "floor", "location", "position", "section", "observation", "item")
    for field in text_fields:
        if field in data:
            product[field] = str(data[field] or "").strip()
    if "description" in data:
        product["name"] = product["description"]
    elif "name" in data:
        product["description"] = product["name"]
    if not product.get("name"):
        raise ValueError("La descripción del producto es obligatoria.")
    if "group" in data:
        product["category"] = product["group"]
        product["source_sheet"] = product["group"]
    if "stock" in data:
        stock = int(data["stock"])
        stock_mode = data.get("stock_mode", "set")
        if stock < 0 or stock_mode not in ("set", "add"):
            raise ValueError("El stock no puede ser negativo.")
        if stock_mode == "add":
            if stock == 0:
                raise ValueError("La cantidad a agregar debe ser mayor que cero.")
            product["stock"] += stock
            movements.append({
                "id": max((int(movement["id"]) for movement in movements), default=0) + 1,
                "product_id": product["id"],
                "movement_type": "Entrada",
                "quantity": stock,
                "floor": product.get("floor", ""),
                "location": product.get("location", ""),
                "note": "Cantidad agregada desde Inventario",
                "created_at": datetime.now().isoformat(timespec="seconds"),
            })
        else:
            product["stock"] = stock
    if "min_stock" in data:
        product["min_stock"] = int(data["min_stock"])
        if product["min_stock"] < 0:
            raise ValueError("El stock minimo no puede ser negativo.")
    if "price" in data:
        product["price"] = float(data["price"])
        if product["price"] < 0:
            raise ValueError("El precio no puede ser negativo.")
    product["sku"] = sku
    product["updated_at"] = datetime.now().isoformat(timespec="seconds")
    write_data(products, movements)
    return True, "Producto actualizado correctamente."


def delete_product(product_id):
    products, movements = read_data()
    product = next((item for item in products if item["id"] == int(product_id)), None)
    if not product:
        return False, "Producto no encontrado."
    identity = _product_identity(product)
    products = [item for item in products if _product_identity(item) != identity]
    write_data(products, movements)
    return True, "Producto eliminado correctamente."


def add_movement(data):
    product_id = int(data["product_id"])
    quantity = int(data["quantity"])
    movement_type = data["movement_type"]
    group = _normalize_movement_group(data.get("group", ""))
    submitted_location = str(data.get("location", "")).strip()
    floor = str(data.get("floor") or (submitted_location if submitted_location in ("Piso 1", "Piso 5") else "Piso 1")).strip()
    if quantity <= 0 or movement_type not in ("Entrada", "Salida") or group not in ("VENTA", "PROYECTO", "ACTIVO", "SERVICIO") or floor not in ("Piso 1", "Piso 5"):
        raise ValueError("Movimiento no valido.")
    products, movements = read_data()
    product = _resolve_movement_product(products, product_id, group, movement_type)
    if not product:
        if not any(int(item["id"]) == product_id for item in products):
            return False, "Producto no encontrado."
        return False, "No hay producto registrado en el grupo seleccionado."
    location = str(data.get("location", product.get("location", ""))).strip()
    position = str(data.get("position", product.get("position", ""))).strip()
    if "floor" not in data and submitted_location in ("Piso 1", "Piso 5"):
        location = str(product.get("location", "")).strip()
    product["floor"] = floor
    product["location"] = location
    product["position"] = position
    for field in ("description", "part_number", "model", "project"):
        value = str(data.get(field, "")).strip()
        if value:
            product[field] = value
            if field == "description":
                product["name"] = value
    new_stock = product["stock"] + quantity if movement_type == "Entrada" else product["stock"] - quantity
    if new_stock < 0:
        return False, "No hay stock suficiente para esta salida."
    product["stock"] = new_stock
    product["updated_at"] = datetime.now().isoformat(timespec="seconds")
    movements.append({
        "id": max((movement["id"] for movement in movements), default=0) + 1,
        "product_id": int(product["id"]),
        "movement_type": movement_type,
        "quantity": quantity,
        "floor": floor,
        "location": location,
        "position": position,
        "project": str(data.get("project", "")).strip(),
        "group": group,
        "note": str(data.get("note", "")).strip(),
        "created_at": str(data.get("movement_date") or datetime.now().date().isoformat()),
    })
    write_data(products, movements)
    return True, "Movimiento registrado correctamente."


def update_movement(movement_id, data):
    try:
        product_id = int(data["product_id"])
        quantity = int(data["quantity"])
    except (KeyError, TypeError, ValueError):
        raise ValueError("Movimiento no valido.")
    movement_type = str(data.get("movement_type", "")).strip()
    group = _normalize_movement_group(data.get("group", ""))
    submitted_location = str(data.get("location", "")).strip()
    floor = str(data.get("floor") or (submitted_location if submitted_location in ("Piso 1", "Piso 5") else "Piso 1")).strip()
    movement_date = str(data.get("movement_date", "")).strip()
    if quantity <= 0 or movement_type not in ("Entrada", "Salida") or group not in ("VENTA", "PROYECTO", "ACTIVO", "SERVICIO") or floor not in ("Piso 1", "Piso 5") or not movement_date:
        raise ValueError("Movimiento no valido.")

    products, movements = read_data()
    movement = next((item for item in movements if int(item["id"]) == int(movement_id)), None)
    if not movement:
        return False, "Movimiento no encontrado."
    old_product = next((item for item in products if int(item["id"]) == int(movement["product_id"])), None)
    if not old_product:
        return False, "Producto no encontrado."
    new_product = _resolve_movement_product(products, product_id, group, movement_type)
    if not new_product:
        if not any(int(item["id"]) == product_id for item in products):
            return False, "Producto no encontrado."
        return False, "No hay producto registrado en el grupo seleccionado."
    location = str(data.get("location", new_product.get("location", ""))).strip()
    position = str(data.get("position", new_product.get("position", ""))).strip()
    if "floor" not in data and submitted_location in ("Piso 1", "Piso 5"):
        location = str(new_product.get("location", "")).strip()
    new_product["floor"] = floor
    new_product["location"] = location
    new_product["position"] = position

    old_quantity = int(movement["quantity"])
    old_effect = old_quantity if movement["movement_type"] == "Entrada" else -old_quantity
    new_effect = quantity if movement_type == "Entrada" else -quantity
    stock_after_reversal = int(old_product["stock"]) - old_effect
    if stock_after_reversal < 0:
        return False, "No se puede editar: el stock de este movimiento ya fue utilizado."
    if old_product is new_product:
        stock_after_update = stock_after_reversal + new_effect
    else:
        stock_after_update = int(new_product["stock"]) + new_effect
    if stock_after_update < 0:
        return False, "No hay stock suficiente para esta salida."

    if old_product is not new_product:
        old_product["stock"] = stock_after_reversal
        new_product["stock"] = int(new_product["stock"]) + new_effect
        new_product["updated_at"] = datetime.now().isoformat(timespec="seconds")
    else:
        old_product["stock"] = stock_after_update
    old_product["updated_at"] = datetime.now().isoformat(timespec="seconds")
    movement.update({
        "product_id": int(new_product["id"]),
        "movement_type": movement_type,
        "quantity": quantity,
        "floor": floor,
        "location": location,
        "position": position,
        "project": str(data.get("project", "")).strip(),
        "group": group,
        "note": str(data.get("note", "")).strip(),
        "created_at": movement_date,
    })
    write_data(products, movements)
    return True, "Movimiento actualizado correctamente."


def delete_movement(movement_id):
    products, movements = read_data()
    movement = next((item for item in movements if item["id"] == int(movement_id)), None)
    if not movement:
        return False, "Movimiento no encontrado."
    product = next((item for item in products if item["id"] == movement["product_id"]), None)
    if product:
        quantity = int(movement["quantity"])
        if movement["movement_type"] == "Entrada":
            product["stock"] = max(0, int(product["stock"]) - quantity)
        else:
            product["stock"] += quantity
        product["updated_at"] = datetime.now().isoformat(timespec="seconds")
    movements.remove(movement)
    write_data(products, movements)
    return True, "Movimiento eliminado correctamente."

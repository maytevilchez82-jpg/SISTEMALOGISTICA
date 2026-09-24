import os
from datetime import datetime
from io import BytesIO
from pathlib import Path
from threading import Lock

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from pyxlsb import open_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_EXCEL_PATH = BASE_DIR / "inventario.xlsx"
SOURCE_XLSB_PATH = BASE_DIR / "INVENTARIO DE EQUIPOS ENERO 2026.xlsb"


def _resolve_excel_path():
    env_path = os.getenv("EXCEL_PATH") or os.getenv("DB_PATH")
    if env_path:
        return Path(env_path).expanduser().resolve()

    custom_excel = (
        sorted(BASE_DIR.glob("*.xlsb"))
        + sorted(BASE_DIR.glob("*.xlsx"))
        + sorted(BASE_DIR.glob("*.xlsm"))
        + sorted(BASE_DIR.glob("*.xls"))
    )
    if custom_excel:
        return custom_excel[0]

    return DEFAULT_EXCEL_PATH


SOURCE_EXCEL_PATH = _resolve_excel_path()
EXCEL_PATH = DEFAULT_EXCEL_PATH if SOURCE_EXCEL_PATH.suffix.lower() == ".xlsb" else SOURCE_EXCEL_PATH
excel_lock = Lock()

PRODUCT_HEADERS = ["id", "sku", "name", "description", "brand", "model", "part_number", "unit", "group", "category", "floor", "location", "position", "observation", "source_sheet", "item", "stock", "min_stock", "price", "updated_at"]
MOVEMENT_HEADERS = ["id", "product_id", "movement_type", "quantity", "floor", "location", "note", "created_at"]


def _rows_from_sheet(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row)]


def read_data():
    with excel_lock:
        workbook = load_workbook(EXCEL_PATH, data_only=True)
        products = _rows_from_sheet(workbook["Productos"])
        movements = _rows_from_sheet(workbook["Movimientos"])
    for product in products:
        product.setdefault("unit", "Unidades")
        product.setdefault("description", product.get("name", ""))
        product.setdefault("brand", "")
        product.setdefault("model", "")
        product.setdefault("part_number", "")
        product.setdefault("group", product.get("category", "Consumibles"))
        product.setdefault("category", product.get("group", "Consumibles"))
        product.setdefault("floor", "")
        product.setdefault("location", "")
        product.setdefault("position", "")
        product.setdefault("observation", "")
        product.setdefault("source_sheet", "")
        product.setdefault("item", product.get("id", ""))
    return products, movements


def write_data(products, movements):
    workbook = Workbook()
    products_sheet = workbook.active
    products_sheet.title = "Productos"
    movements_sheet = workbook.create_sheet("Movimientos")
    products_sheet.append(PRODUCT_HEADERS)
    movements_sheet.append(MOVEMENT_HEADERS)
    for product in products:
        products_sheet.append([product.get(header) for header in PRODUCT_HEADERS])
    for movement in movements:
        movements_sheet.append([movement.get(header) for header in MOVEMENT_HEADERS])
    for sheet in (products_sheet, movements_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    products_sheet.column_dimensions["B"].width = 16
    products_sheet.column_dimensions["C"].width = 28
    products_sheet.column_dimensions["D"].width = 18
    movements_sheet.column_dimensions["E"].width = 28
    with excel_lock:
        workbook.save(EXCEL_PATH)


def initialize_database():
    if EXCEL_PATH.exists():
        return
    if SOURCE_EXCEL_PATH.exists() and SOURCE_EXCEL_PATH.suffix.lower() == ".xlsb":
        write_data(_import_xlsb_products(SOURCE_EXCEL_PATH), [])
        return
    if SOURCE_XLSB_PATH.exists():
        write_data(_import_xlsb_products(SOURCE_XLSB_PATH), [])
        return
    now = datetime.now().isoformat(timespec="seconds")
    products = [
        {"id": 1, "sku": "TEC-001", "name": "Teclado mecanico", "unit": "Cajas", "group": "Perifericos", "category": "Perifericos", "stock": 24, "min_stock": 8, "price": 49.90, "updated_at": now},
        {"id": 2, "sku": "MON-024", "name": "Monitor 24 pulgadas", "unit": "Cajas", "group": "Monitores", "category": "Monitores", "stock": 8, "min_stock": 4, "price": 189.00, "updated_at": now},
        {"id": 3, "sku": "CAB-110", "name": "Cable USB-C reforzado", "unit": "Cajas", "group": "Accesorios", "category": "Accesorios", "stock": 42, "min_stock": 12, "price": 14.50, "updated_at": now},
        {"id": 4, "sku": "AUD-301", "name": "Audifonos estudio", "unit": "Cajas", "group": "Audio", "category": "Audio", "stock": 3, "min_stock": 5, "price": 79.90, "updated_at": now},
        {"id": 5, "sku": "MOU-089", "name": "Mouse ergonomico", "unit": "Cajas", "group": "Perifericos", "category": "Perifericos", "stock": 17, "min_stock": 6, "price": 32.00, "updated_at": now},
    ]
    movements = [
        {"id": 1, "product_id": 1, "movement_type": "Entrada", "quantity": 10, "note": "Compra proveedor", "created_at": now},
        {"id": 2, "product_id": 4, "movement_type": "Salida", "quantity": 2, "note": "Pedido #1048", "created_at": now},
        {"id": 3, "product_id": 3, "movement_type": "Entrada", "quantity": 20, "note": "Reposicion", "created_at": now},
        {"id": 4, "product_id": 2, "movement_type": "Salida", "quantity": 1, "note": "Pedido #1047", "created_at": now},
    ]
    write_data(products, movements)


def _import_xlsb_products(source_path=None):
    source_path = Path(source_path) if source_path else SOURCE_XLSB_PATH
    products = []
    next_id = 1
    with open_workbook(source_path) as workbook:
        for sheet_name in workbook.sheets:
            with workbook.get_sheet(sheet_name) as sheet:
                rows = ([cell.v for cell in row] for row in sheet.rows())
                headers = None
                for row in rows:
                    normalized = [str(value).strip().upper() if value is not None else "" for value in row]
                    if "ITEM" in normalized and "DESCRIPCIÓN" in normalized:
                        headers = normalized
                        break
                if not headers:
                    continue
                for row in rows:
                    values = list(row)
                    if not any(value is not None for value in values):
                        continue
                    data = {headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))}
                    description = str(data.get("DESCRIPCIÓN") or "").strip()
                    if not description:
                        continue
                    item = data.get("ITEM")
                    item_text = str(int(item)) if isinstance(item, float) and item.is_integer() else str(item or "")
                    quantity = data.get("CANTIDAD") or 0
                    location = data.get("CASILLERO")
                    location_text = f"Locker {int(location)}" if isinstance(location, float) and location.is_integer() else str(location or "").strip()
                    products.append({
                        "id": next_id,
                        "sku": f"{sheet_name}-{item_text}",
                        "name": description,
                        "description": description,
                        "brand": str(data.get("MARCA") or "").strip(),
                        "model": str(data.get("MODELO") or "").strip(),
                        "part_number": str(data.get("PART NUMBER") or "").strip(),
                        "unit": "Unidades",
                        "group": sheet_name,
                        "category": sheet_name,
                        "floor": "",
                        "location": location_text,
                        "position": str(data.get("POSICION") or "").strip(),
                        "observation": str(data.get("OBSERVACION") or "").strip(),
                        "source_sheet": sheet_name,
                        "item": item_text,
                        "stock": int(quantity),
                        "min_stock": 0,
                        "price": 0,
                        "updated_at": datetime.now().isoformat(timespec="seconds"),
                    })
                    next_id += 1
    return products


def get_products():
    products, _ = read_data()
    return sorted(products, key=lambda item: item["name"])


def export_products_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventario"
    logo_path = BASE_DIR / "static" / "NAKAMA LOGO 1.jpeg"
    if logo_path.exists():
        logo = Image(logo_path)
        logo.width = 200
        logo.height = 75
        sheet.add_image(logo, "A1")
    sheet.merge_cells("B1:N2")
    title = sheet["B1"]
    title.value = "INVENTARIO DE NAKAMA SOLUCIONES"
    title.font = Font(name="Calibri", size=18, bold=True, color="18236F")
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells("B3:N3")
    download_date = sheet["B3"]
    download_date.value = f"Fecha de descarga: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    download_date.font = Font(name="Calibri", size=11, italic=True, color="5C747C")
    download_date.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 40
    sheet.row_dimensions[2].height = 40
    sheet.row_dimensions[3].height = 10

    headers = ["ID", "HOJA", "ITEM", "PRODUCTO", "MARCA", "MODELO", "PART NUMBER", "UNIDAD", "CANTIDAD", "PISO", "LOCKER", "POSICION", "OBSERVACION", "ESTADO"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=4, column=column, value=header)
    for product in get_products():
        sheet.append([
            product["id"],
            product.get("source_sheet", ""),
            product.get("item", ""),
            product["name"],
            product.get("brand", ""),
            product.get("model", ""),
            product.get("part_number", ""),
            product.get("unit", "Unidades"),
            product["stock"],
            product.get("floor", ""),
            product.get("location", ""),
            product.get("position", ""),
            product.get("observation", ""),
            "Reponer" if product["stock"] <= product["min_stock"] else "Saludable",
        ])

    last_row = max(sheet.max_row, 5)
    table = Table(displayName="InventarioTabla", ref=f"A4:N{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "A5"
    sheet.column_dimensions["A"].width = 12
    for column in "BCDEFGHIJKLMN":
        sheet.column_dimensions[column].width = 20
    sheet.column_dimensions["D"].width = 36
    sheet.column_dimensions["K"].width = 16
    sheet.column_dimensions["M"].width = 32

    movements_sheet = workbook.create_sheet("Movimientos")
    movement_headers = ["ID", "PRODUCTO ID", "CODIGO", "PRODUCTO", "TIPO", "CANTIDAD", "PISO", "LOCKER", "NOTA", "FECHA"]
    movements_sheet.append(movement_headers)
    for movement in get_movements():
        movements_sheet.append([
            movement.get("id"),
            movement.get("product_id"),
            movement.get("sku", ""),
            movement.get("product_name", ""),
            movement.get("movement_type", ""),
            movement.get("quantity", 0),
            movement.get("floor", ""),
            movement.get("location", ""),
            movement.get("note", ""),
            movement.get("created_at", ""),
        ])
    movements_sheet.freeze_panes = "A2"
    movements_sheet.auto_filter.ref = movements_sheet.dimensions
    for column in "ABCDEFGHIJ":
        movements_sheet.column_dimensions[column].width = 18
    movements_sheet.column_dimensions["D"].width = 32
    movements_sheet.column_dimensions["I"].width = 32

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def get_movements():
    products, movements = read_data()
    products_by_id = {product["id"]: product for product in products}
    return [movement_with_product(movement, products_by_id) for movement in sorted(movements, key=lambda item: item["id"], reverse=True)]


def movement_with_product(movement, products_by_id):
    product = products_by_id.get(movement["product_id"], {})
    return {**movement, "product_name": product.get("name", "Producto eliminado"), "sku": product.get("sku", "")}


def add_product(data):
    products, movements = read_data()
    sku = data["sku"].strip().upper()
    if any(product["sku"] == sku for product in products):
        return False, "El SKU ya existe."
    products.append({
        "id": max((product["id"] for product in products), default=0) + 1,
        "sku": sku,
        "name": data["name"].strip(),
        "brand": str(data.get("brand", "")).strip(),
        "unit": data.get("unit", "Unidades").strip(),
        "group": data["group"].strip(),
        "category": data["group"].strip(),
        "floor": str(data.get("floor", "Piso 1")).strip(),
        "location": str(data.get("location", "Locker 1")).strip(),
        "stock": int(data.get("stock", 0)),
        "min_stock": int(data.get("min_stock", 5)),
        "price": float(data.get("price", 0)),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    })
    write_data(products, movements)
    return True, "Producto creado correctamente."


def add_movement(data):
    product_id = int(data["product_id"])
    quantity = int(data["quantity"])
    movement_type = data["movement_type"]
    floor = str(data.get("floor", "Piso 1")).strip()
    location = str(data.get("location", "Locker 1")).strip()
    if quantity <= 0 or movement_type not in ("Entrada", "Salida") or floor not in ("Piso 1", "Piso 5") or not location:
        raise ValueError("Movimiento no valido.")
    products, movements = read_data()
    product = next((item for item in products if item["id"] == product_id), None)
    if not product:
        return False, "Producto no encontrado."
    new_stock = product["stock"] + quantity if movement_type == "Entrada" else product["stock"] - quantity
    if new_stock < 0:
        return False, "No hay stock suficiente para esta salida."
    product["stock"] = new_stock
    product["updated_at"] = datetime.now().isoformat(timespec="seconds")
    movements.append({
        "id": max((movement["id"] for movement in movements), default=0) + 1,
        "product_id": product_id,
        "movement_type": movement_type,
        "quantity": quantity,
        "floor": floor,
        "location": location,
        "note": str(data.get("note", "")).strip(),
        "created_at": str(data.get("movement_date") or datetime.now().date().isoformat()),
    })
    write_data(products, movements)
    return True, "Movimiento registrado correctamente."
